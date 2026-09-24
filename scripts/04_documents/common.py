# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

def newdoc():
    doc=Document()
    st=doc.styles['Normal']
    st.font.name='Times New Roman'; st.font.size=Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    pf=st.paragraph_format
    pf.line_spacing_rule=WD_LINE_SPACING.DOUBLE
    pf.space_after=Pt(0)
    for s in doc.sections:
        s.left_margin=s.right_margin=Inches(1.0); s.top_margin=s.bottom_margin=Inches(1.0)
    return doc

def mk(doc):
    def P(text='',align=None,italic=False,bold=False,indent=True,space_before=0,size=None):
        p=doc.add_paragraph()
        p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.DOUBLE
        p.paragraph_format.space_before=Pt(space_before)
        p.paragraph_format.space_after=Pt(0)
        if indent: p.paragraph_format.first_line_indent=Inches(0.5)
        if align: p.alignment=align
        r=p.add_run(text); r.italic=italic; r.bold=bold
        if size: r.font.size=Pt(size)
        return p
    def RICH(parts, align=None, indent=True, space_before=0):
        p=doc.add_paragraph()
        p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.DOUBLE
        p.paragraph_format.space_before=Pt(space_before)
        p.paragraph_format.space_after=Pt(0)
        if indent: p.paragraph_format.first_line_indent=Inches(0.5)
        if align: p.alignment=align
        for item in parts:
            t,it=item[0],item[1]
            r=p.add_run(t); r.italic=it
            if len(item)>2 and item[2]: r.bold=True
        return p
    def H(text,level=1):
        p=doc.add_paragraph()
        p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.DOUBLE
        p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(0)
        r=p.add_run(text); r.bold=True; r.font.size=Pt(12)
        if level>2: r.bold=False; r.italic=True
        return p
    def HANG(text):
        p=doc.add_paragraph()
        p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.DOUBLE
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.left_indent=Inches(0.5)
        p.paragraph_format.first_line_indent=Inches(-0.5)
        p.add_run(text)
        return p
    def TABLE(rows, header=True):
        t=doc.add_table(rows=1,cols=len(rows[0])); t.style='Table Grid'
        for i,h in enumerate(rows[0]):
            c=t.rows[0].cells[i]; c.text=str(h)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
                for r in p.runs: r.bold=header; r.font.size=Pt(10)
        for row in rows[1:]:
            cells=t.add_row().cells
            for i,v in enumerate(row):
                cells[i].text=str(v)
                for p in cells[i].paragraphs:
                    p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE
                    for r in p.runs: r.font.size=Pt(10)
        return t
    return P,RICH,H,HANG,TABLE
